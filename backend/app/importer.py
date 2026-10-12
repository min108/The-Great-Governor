import re
from openpyxl import load_workbook

from .db import ROOT

DATA_XLSX = ROOT / "sources" / "data_standards.xlsx"
ALGO_XLSX = ROOT / "sources" / "algorithms.xlsx"
CODE_RE = re.compile(r"D\.\d{2}\.\d{2}\.\d{3}")


def _cell(value):
    if value is None:
        return ""
    return str(value).strip()


def import_standards(conn):
    wb = load_workbook(DATA_XLSX, read_only=True, data_only=True)
    rows = []
    for sheet in wb.sheetnames:
        if sheet in ("编制说明", "总览索引"):
            continue
        ws = wb[sheet]
        header = None
        domain_code = sheet[:2]
        domain_name = sheet[3:].strip() if len(sheet) > 3 else sheet
        for i, row in enumerate(ws.iter_rows(values_only=True)):
            values = [_cell(c) for c in row]
            if i == 0:
                header = values
                continue
            if not values or not values[0].startswith("D."):
                continue
            padded = values + [""] * 14
            rows.append(
                (
                    padded[0],
                    domain_code,
                    domain_name,
                    padded[1],
                    padded[2],
                    padded[3],
                    padded[4],
                    padded[5],
                    padded[6],
                    padded[7],
                    padded[8],
                    padded[9],
                    padded[10],
                    padded[11],
                    padded[12],
                    padded[13],
                )
            )
    wb.close()
    conn.executemany(
        """INSERT OR REPLACE INTO data_standards
        (code, domain_code, domain_name, category2, category3, name, form, asset_class,
         meaning, reference_form, constraints, basis, positive_example, negative_example,
         sensitivity, related_rules)
        VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)""",
        rows,
    )
    return len(rows)


def import_algorithms(conn):
    wb = load_workbook(ALGO_XLSX, read_only=True, data_only=True)
    ws = wb["总表"]
    rows = []
    for i, row in enumerate(ws.iter_rows(values_only=True)):
        if i == 0:
            continue
        values = [_cell(c) for c in row]
        if not values or not values[0].startswith("A."):
            continue
        padded = values + [""] * 11
        rows.append(tuple(padded[:11]))
    wb.close()
    conn.executemany(
        """INSERT OR REPLACE INTO algorithms
        (code, domain, subcategory, algo_type, name, scenario, data_needs, method,
         judgment, output_template, related_rules)
        VALUES (?,?,?,?,?,?,?,?,?,?,?)""",
        rows,
    )
    return len(rows)


def link_algorithm_data(conn):
    existing = {r[0] for r in conn.execute("SELECT code FROM data_standards")}
    links = []
    n = 0
    for code, needs in conn.execute("SELECT code, data_needs FROM algorithms"):
        seen = []
        for match in CODE_RE.findall(needs or ""):
            if match in existing and match not in seen:
                seen.append(match)
        for match in seen:
            n += 1
            links.append(
                (
                    f"rel-algo-data-{n}",
                    "algorithm",
                    code,
                    "standard",
                    match,
                    "引用数据",
                    "清单编制口径",
                    "分析该思路时需要核对该数据项",
                    "材料写明",
                    "监督思路清单「数据需求」栏已列明该项",
                )
            )
    conn.executemany(
        """INSERT OR REPLACE INTO relations
        (id, from_kind, from_id, to_kind, to_id, rel_type, period, condition, confirm_status, evidence)
        VALUES (?,?,?,?,?,?,?,?,?,?)""",
        links,
    )
    return len(links)
